"""Create a human-review workbook comparing two OCR engine outputs."""
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

KEYS = ("pdf_page", "grid_row", "grid_col")
FIELDS = ("serial_no", "voter_id", "name", "relation_type", "relation_name", "house_no", "age", "gender")


def _read_rows(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Voters"]
    headers = next(ws.values)
    rows = {tuple(row[headers.index(key)] for key in KEYS): dict(zip(headers, row))
            for row in ws.iter_rows(min_row=2, values_only=True)}
    wb.close()
    return rows


def compare_workbooks(baseline_path, candidate_path, output_path):
    baseline = _read_rows(baseline_path)
    candidate = _read_rows(candidate_path)
    keys = sorted(baseline.keys() | candidate.keys())
    wb = Workbook()
    summary = wb.active
    summary.title = "Comparison Summary"
    summary.append(["Field", "Tesseract populated", "RapidOCR populated", "Both populated", "Agree", "Disagree", "Only Tesseract", "Only RapidOCR"])
    review = wb.create_sheet("Review Differences")
    review.append(["PDF page", "Grid row", "Grid column", "Field", "Tesseract value", "RapidOCR value", "Difference type"])
    counts = {field: {key: 0 for key in ("base", "rapid", "both", "agree", "disagree", "base_only", "rapid_only")} for field in FIELDS}
    for key in keys:
        left, right = baseline.get(key, {}), candidate.get(key, {})
        for field in FIELDS:
            a, b = left.get(field), right.get(field)
            has_a, has_b = a not in (None, ""), b not in (None, "")
            c = counts[field]
            c["base"] += has_a
            c["rapid"] += has_b
            if has_a and has_b:
                c["both"] += 1
                same = str(a).strip().casefold() == str(b).strip().casefold()
                c["agree" if same else "disagree"] += 1
                if not same:
                    review.append([*key, field, a, b, "Different values"])
            elif has_a:
                c["base_only"] += 1
                review.append([*key, field, a, b, "Only Tesseract populated"])
            elif has_b:
                c["rapid_only"] += 1
                review.append([*key, field, a, b, "Only RapidOCR populated"])
    for field, c in counts.items():
        summary.append([field, c["base"], c["rapid"], c["both"], c["agree"], c["disagree"], c["base_only"], c["rapid_only"]])
    for ws in (summary, review):
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for col in ws.columns:
            width = min(52, max(12, max(len(str(cell.value or "")) for cell in list(col)[:80]) + 2))
            ws.column_dimensions[get_column_letter(col[0].column)].width = width
    summary.cell(row=summary.max_row + 2, column=1, value="Note: These are OCR-to-OCR comparisons, not accuracy scores. Review differences against the source PDF.")
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return len(keys), review.max_row - 1


def build_dual_outputs(tesseract_path, rapidocr_path, unfiltered_path, consensus_path):
    """Write unfiltered engine reads and a reviewable best-available file.

    RapidOCR is the preferred candidate source due to its higher field
    coverage. Conflicts remain flagged, and consensus-only values are included
    separately. Neither output represents ground truth.
    """
    left = _read_rows(tesseract_path)
    right = _read_rows(rapidocr_path)
    keys = sorted(left.keys() | right.keys())
    left_wb = load_workbook(tesseract_path, read_only=True, data_only=True)
    right_wb = load_workbook(rapidocr_path, read_only=True, data_only=True)
    left_headers = list(next(left_wb["Voters"].values))
    right_headers = list(next(right_wb["Voters"].values))
    left_wb.close(); right_wb.close()

    full = Workbook()
    all_ws = full.active
    all_ws.title = "All Engine Reads"
    extras = [h for h in left_headers if h not in KEYS]
    extras += [h for h in right_headers if h not in KEYS and h not in extras]
    all_ws.append([*KEYS, *[f"Tesseract | {h}" for h in extras], *[f"RapidOCR | {h}" for h in extras]])
    for key in keys:
        a, b = left.get(key, {}), right.get(key, {})
        all_ws.append([*key, *[a.get(h) for h in extras], *[b.get(h) for h in extras]])

    settled = Workbook()
    data_ws = settled.active
    data_ws.title = "Best Available Candidates"
    consensus_ws = settled.create_sheet("Consensus Only")
    review_ws = settled.create_sheet("Review Queue")
    fields = [f for f in FIELDS if f in left_headers or f in right_headers]
    data_ws.append([*KEYS, *sum(([f, f"{f}_status"] for f in fields), [])])
    consensus_ws.append([*KEYS, *fields])
    review_ws.append([*KEYS, "Field", "Tesseract value", "RapidOCR value", "Reason"])
    for key in keys:
        a, b = left.get(key, {}), right.get(key, {})
        values = [*key]
        agreed_values = [*key]
        for field in fields:
            av, bv = a.get(field), b.get(field)
            has_a, has_b = av not in (None, ""), bv not in (None, "")
            if has_a and has_b and str(av).strip().casefold() == str(bv).strip().casefold():
                values.extend([bv, "AGREED_UNVERIFIED"])
                agreed_values.append(bv)
            elif has_a or has_b:
                selected, status = (bv, "RAPIDOCR_PREFERRED_REVIEW") if has_b else (av, "TESSERACT_FALLBACK_REVIEW")
                values.extend([selected, status])
                agreed_values.append(None)
                review_ws.append([*key, field, av, bv, "Engines disagree" if has_a and has_b else "Only one engine populated this field"])
            else:
                values.extend([None, "MISSING_BOTH"])
                agreed_values.append(None)
                review_ws.append([*key, field, av, bv, "Neither engine populated this field"])
        data_ws.append(values)
        consensus_ws.append(agreed_values)

    method_ws = settled.create_sheet("Read Me")
    method_ws.append(["Best Available Candidates prefers RapidOCR due to its higher field coverage; Tesseract fills only RapidOCR blanks."])
    method_ws.append(["RAPIDOCR_PREFERRED_REVIEW, TESSERACT_FALLBACK_REVIEW and missing statuses require source-PDF review."])
    method_ws.append(["Consensus Only contains values both engines read identically. Agreement is not ground-truth verification."])
    method_ws.append(["Review Queue lists disagreements, one-engine-only values and values missing from both engines."])
    for wb in (full, settled):
        for ws in wb.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
            for col in ws.columns:
                width = min(52, max(12, max(len(str(cell.value or "")) for cell in list(col)[:80]) + 2))
                ws.column_dimensions[get_column_letter(col[0].column)].width = width
    Path(unfiltered_path).parent.mkdir(parents=True, exist_ok=True)
    Path(consensus_path).parent.mkdir(parents=True, exist_ok=True)
    full.save(unfiltered_path)
    settled.save(consensus_path)
    return len(keys), review_ws.max_row - 1
