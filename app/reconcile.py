from collections import Counter
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter

def reconcile(rows,targets):
    serials=[str(r['serial_no']) for r in rows if r.get('serial_no') is not None]; sc=Counter(serials)
    duplicates=sorted(x for x,n in sc.items() if n>1); start=targets.get('target_start_serial'); end=targets.get('target_end_serial')
    expected=set(map(str,range(start,end+1))) if start is not None and end is not None and end>=start else set()
    missing=sorted(expected-set(serials),key=int) if expected else []
    out_of_range=sorted(x for x in set(serials) if expected and x not in expected)
    ids=[r['voter_id'] for r in rows if r.get('voter_id')]; ic=Counter(ids); dup_ids=sorted(x for x,n in ic.items() if n>1)
    genders=Counter((r.get('gender') or '').strip().lower() for r in rows)
    checks=[('Total electors',targets.get('target_total'),len(rows)),('Male electors',targets.get('target_male'),genders['male']),('Female electors',targets.get('target_female'),genders['female']),('Third-gender electors',targets.get('target_third_gender'),genders['third gender'])]
    report=[]
    for label,target,actual in checks:
        status='TARGET_EXTRACTION_FAILED' if label=='Total electors' and target is None else ('NOT_CHECKED' if target is None else ('PASS' if target==actual else 'REVIEW_REQUIRED'))
        report.append({'check':label,'target':target,'actual':actual,'difference':None if target is None else actual-target,'status':status})
    for label,issues in [('Duplicate serial numbers',duplicates),('Missing serial numbers',missing),('Out-of-range serial numbers',out_of_range),('Duplicate voter IDs',dup_ids)]:
        unchecked=(not rows) or (label=='Missing serial numbers' and not expected)
        report.append({'check':label,'target':0,'actual':len(issues),'difference':len(issues),'status':'NOT_CHECKED' if unchecked else ('PASS' if not issues else 'REVIEW_REQUIRED')})
    missing_fields={k:sum(not r.get(k) for r in rows) for k in ('serial_no','voter_id','name','age','gender')}
    report.append({'check':'Missing voter IDs','target':0,'actual':missing_fields['voter_id'],'difference':missing_fields['voter_id'],'status':'NOT_CHECKED' if not rows else ('PASS' if not missing_fields['voter_id'] else 'REVIEW_REQUIRED')})
    for k,v in missing_fields.items():report.append({'check':f'Missing {k}','target':0,'actual':v,'difference':v,'status':'NOT_CHECKED' if not rows else ('PASS' if not v else 'REVIEW_REQUIRED')})
    return report,duplicates,missing,dup_ids,out_of_range

def write_workbook(out_path,rows,targets,report,duplicates,missing,dup_ids=(),diagnostics=(),out_of_range=()):
    wb=Workbook(); ws=wb.active;ws.title='Voters'
    fields=['pdf_page','grid_row','grid_col','card_box_px','serial_no','serial_status','serial_candidates','voter_id','voter_id_status','voter_id_selection','voter_id_candidates','voter_id_raw','name','relation_type','relation_name','house_no','age','gender','gender_candidates','gender_raw','raw_text']
    ws.append(fields)
    for r in rows:ws.append([r.get(k) for k in fields])
    q=wb.create_sheet('Reconciliation Report');q.append(['Check','Target','Actual','Difference','Status'])
    for r in report:q.append([r['check'],r['target'],r['actual'],r['difference'],r['status']])
    sr=wb.create_sheet('Serial Review');sr.append(['Issue','Value','PDF page','Grid row','Grid column'])
    for issue,values in [('Duplicate',duplicates),('Missing from expected range',missing)]:
        for value in values:
            matched=[r for r in rows if str(r.get('serial_no'))==str(value)]
            if matched:
                for r in matched:sr.append([issue,value,r.get('pdf_page'),r.get('grid_row'),r.get('grid_col')])
            else:sr.append([issue,value])
    for value in out_of_range:
        for r in rows:
            if str(r.get('serial_no'))==str(value):
                sr.append(['Out of expected range',value,r.get('pdf_page'),r.get('grid_row'),r.get('grid_col')])
    for r in rows:
        if r.get('serial_status')=='AMBIGUOUS':
            sr.append(['Ambiguous OCR candidates',r.get('serial_candidates'),r.get('pdf_page'),r.get('grid_row'),r.get('grid_col')])
        elif r.get('serial_status')=='CANDIDATE_REQUIRES_VERIFICATION':
            sr.append(['Candidate requires visual verification',r.get('serial_candidates'),r.get('pdf_page'),r.get('grid_row'),r.get('grid_col')])
        elif r.get('serial_status')=='NOT_FOUND':
            sr.append(['Printed serial not found',None,r.get('pdf_page'),r.get('grid_row'),r.get('grid_col')])
    rv=wb.create_sheet('Review Required');rv.append(fields)
    for r in rows:
        if (any(not r.get(k) for k in ('serial_no','voter_id','name','age','gender')) or
            r.get('serial_status')!='VERIFIED' or r.get('voter_id_status')!='VERIFIED'):
            rv.append([r.get(k) for k in fields])
    pg=wb.create_sheet('Page Diagnostics');pg.append(['Page','Voter cards accepted','Inked grid slots','Failed field count','Processing error','Duration seconds'])
    for d in diagnostics:pg.append([d['page'],d['detected'],d.get('grid_candidates',d['detected']),d['failed_fields'],d['error'],d['duration_seconds']])
    ev=wb.create_sheet('Target Evidence');ev.append(['Field','Value'])
    for k,v in targets.items():ev.append([k,v if not isinstance(v,dict) else str(v)])
    for sh in wb.worksheets:
        sh.freeze_panes='A2';sh.auto_filter.ref=sh.dimensions
        for c in sh[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='244062');c.alignment=Alignment(wrap_text=True)
        for col in sh.columns:
            letter=get_column_letter(col[0].column);sh.column_dimensions[letter].width=min(55,max(12,max((len(str(c.value or '')) for c in col[:150]),default=12)+2))
    Path(out_path).parent.mkdir(parents=True,exist_ok=True);wb.save(out_path)
