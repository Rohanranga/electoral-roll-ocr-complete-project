"""Local, read-only dashboard for browsing OCR workbooks and source-card crops."""
from __future__ import annotations

import json
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, quote, urlparse

from openpyxl import load_workbook


PAGE = r"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Electoral Roll OCR Review</title>
<style>
:root{color-scheme:light;--ink:#172b4d;--muted:#60718a;--line:#dce5ef;--blue:#1359a6;--pale:#f4f7fb;--green:#e5f7ef;--amber:#fff3d5;--red:#ffebeb}
*{box-sizing:border-box}body{margin:0;background:var(--pale);color:var(--ink);font:15px/1.45 "Segoe UI",Arial,sans-serif}
header{background:linear-gradient(115deg,#092849,#145b9c);color:white;padding:22px max(22px,calc((100% - 1440px)/2));box-shadow:0 3px 14px #0a254022}
h1{font-size:25px;margin:0 0 4px}header p{margin:0;color:#d5e5f8}.wrap{max-width:1440px;margin:auto;padding:20px}
.toolbar{display:flex;gap:12px;align-items:end;flex-wrap:wrap;background:#fff;border:1px solid var(--line);padding:14px;border-radius:12px}
label{display:grid;gap:5px;font-size:12px;font-weight:700;color:var(--muted)}select,input,button{font:inherit;border:1px solid #cbd6e3;border-radius:8px;padding:9px 11px;background:white;color:var(--ink)}select{min-width:320px}input{min-width:250px}button{cursor:pointer;font-weight:650}button:hover{border-color:var(--blue);color:var(--blue)}button:disabled{opacity:.4;cursor:default}
.metrics{display:grid;grid-template-columns:repeat(4,minmax(130px,1fr));gap:12px;margin:14px 0}.metric{background:#fff;border:1px solid var(--line);border-radius:11px;padding:13px 16px}.metric span{display:block;color:var(--muted);font-size:12px;font-weight:650}.metric strong{display:block;margin-top:4px;font-size:22px}
.layout{display:grid;grid-template-columns:minmax(270px,360px) 1fr;gap:14px;align-items:start}.panel{background:#fff;border:1px solid var(--line);border-radius:12px;overflow:hidden}.panel-head{padding:13px 15px;border-bottom:1px solid var(--line);font-weight:750;display:flex;justify-content:space-between;align-items:center}.record-list{max-height:72vh;overflow:auto}.record-item{width:100%;text-align:left;border:0;border-bottom:1px solid #edf1f6;border-radius:0;padding:11px 14px}.record-item.active{background:#eaf3ff;box-shadow:inset 3px 0 var(--blue)}.record-item b,.record-item small{display:block}.record-item small{color:var(--muted);margin-top:3px}.pill{display:inline-block;padding:2px 7px;border-radius:99px;background:var(--amber);font-size:11px;color:#795400;font-weight:700;margin-left:6px}
.detail{padding:16px}.detail-top{display:flex;justify-content:space-between;gap:10px;align-items:center;margin-bottom:14px}.nav{display:flex;gap:7px}.content{display:grid;grid-template-columns:minmax(300px,1.05fr) minmax(300px,1fr);gap:18px}.image-box{background:#101a27;border-radius:10px;min-height:280px;display:grid;place-items:center;overflow:hidden}.image-box img{max-width:100%;max-height:72vh;object-fit:contain}.empty-image{color:#cfdae8;padding:30px;text-align:center}.caption{font-size:12px;color:var(--muted);margin:7px 2px 0}.fields{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:8px}.field{border:1px solid var(--line);border-radius:8px;padding:9px 10px;min-width:0}.field span{display:block;font-size:11px;color:var(--muted);font-weight:700;text-transform:uppercase;letter-spacing:.035em}.field b{display:block;margin-top:3px;overflow-wrap:anywhere;white-space:pre-wrap}.field.wide{grid-column:1/-1}.evidence{margin-top:14px;border-top:1px solid var(--line);padding-top:10px}.evidence details{margin:7px 0}.evidence summary{cursor:pointer;font-weight:700}.evidence pre{white-space:pre-wrap;overflow-wrap:anywhere;background:#f4f7fa;border-radius:8px;padding:10px;max-height:300px;overflow:auto;font:12px/1.5 Consolas,monospace}.status{padding:16px;color:var(--muted)}.notice{color:#8a3c00;background:#fff3d5;border-radius:8px;padding:10px;margin:10px 0;font-size:13px}
@media(max-width:900px){.layout{grid-template-columns:1fr}.record-list{max-height:35vh}.content{grid-template-columns:1fr}.metrics{grid-template-columns:repeat(2,1fr)}}@media(max-width:560px){.wrap{padding:12px}.metrics{gap:7px}.metric{padding:10px}.metric strong{font-size:19px}select,input{min-width:100%;width:100%}.toolbar label{width:100%}.fields{grid-template-columns:1fr}.field.wide{grid-column:auto}}
</style></head><body>
<header><h1>Electoral Roll OCR Review</h1><p>Browse extracted fields beside each source voter card. OCR values are candidates; verify them against the image.</p></header>
<main class="wrap"><section class="toolbar">
<label>Workbook<select id="workbook"><option>Loading workbooks…</option></select></label>
<label>Search records<input id="search" placeholder="Name, voter ID, serial, house…"></label>
<label>Show<select id="filter"><option value="all">All records</option><option value="review">Needs review</option><option value="missing">Has missing fields</option></select></label>
<button id="refresh">Refresh</button></section>
<section class="metrics" id="metrics"></section>
<div id="message" class="status">Loading…</div>
<section class="layout" id="layout" hidden><aside class="panel"><div class="panel-head"><span>Records</span><span id="record-count"></span></div><div class="record-list" id="record-list"></div></aside>
<article class="panel"><div class="detail" id="detail"><div class="status">Choose a record.</div></div></article></section></main>
<script>
const $=s=>document.querySelector(s);let books=[],records=[],visible=[],selected=0;
const show=v=>v===null||v===undefined||String(v).trim()===""?"—":String(v);
function addMetric(label,value){const d=document.createElement('div');d.className='metric';const a=document.createElement('span'),b=document.createElement('strong');a.textContent=label;b.textContent=value;d.append(a,b);$('#metrics').append(d)}
function setMessage(text,showLayout=false){$('#message').textContent=text;$('#message').hidden=showLayout;$('#layout').hidden=!showLayout}
async function loadBooks(){setMessage('Loading workbooks…');try{const r=await fetch('/api/workbooks');if(!r.ok)throw Error(await r.text());books=await r.json();const s=$('#workbook');s.replaceChildren();if(!books.length){s.add(new Option('No OCR workbooks found',''));setMessage('No *_ocr.xlsx workbooks were found. Run OCR first or choose the folder containing your workbooks when starting the dashboard.');return}for(const b of books)s.add(new Option(`${b.name} (${b.records} records)`,b.name));await loadRecords()}catch(e){setMessage(`Could not load workbooks: ${e.message}`)}}
async function loadRecords(){const name=$('#workbook').value;if(!name)return;setMessage('Reading workbook…');try{const r=await fetch(`/api/records?workbook=${encodeURIComponent(name)}`);if(!r.ok)throw Error(await r.text());const d=await r.json();records=d.records;selected=0;render();}catch(e){setMessage(`Could not read workbook: ${e.message}`)}}
function needsReview(r){return ['serial_status','voter_id_status'].some(k=>show(r[k]).includes('VERIFICATION')||show(r[k])==='AMBIGUOUS'||show(r[k])==='NOT_FOUND')||['serial_no','voter_id','name','age','gender'].some(k=>show(r[k])==='—')}
function missingAny(r){return ['serial_no','voter_id','name','relation_name','house_no','age','gender'].some(k=>show(r[k])==='—')}
function render(){const q=$('#search').value.trim().toLowerCase(),mode=$('#filter').value;visible=records.map((r,i)=>({r,i})).filter(({r})=>{const hay=Object.values(r).join(' ').toLowerCase();return(!q||hay.includes(q))&&(mode==='all'||(mode==='review'?needsReview(r):missingAny(r)))});$('#record-count').textContent=`${visible.length} / ${records.length}`;const list=$('#record-list');list.replaceChildren();visible.forEach(({r,i})=>{const b=document.createElement('button');b.className='record-item'+(i===selected?' active':'');b.onclick=()=>{selected=i;render()};const title=document.createElement('b');title.textContent=`${show(r.serial_no)} · ${show(r.name)}`;const sub=document.createElement('small');sub.textContent=`Page ${show(r.pdf_page)} · Voter ID ${show(r.voter_id)}${needsReview(r)?' · Review':''}`;b.append(title,sub);list.append(b)});if(!visible.length){setMessage(records.length?'No records match these filters.':'This workbook has no voter rows.');return}const current=visible.findIndex(x=>x.i===selected);if(current<0)selected=visible[0].i;drawDetail(records[selected],visible.findIndex(x=>x.i===selected),visible.length);$('#metrics').replaceChildren();const missing=records.filter(missingAny).length,review=records.filter(needsReview).length;addMetric('Voter records',records.length.toLocaleString());addMetric('Needs review',review.toLocaleString());addMetric('Missing fields',missing.toLocaleString());addMetric('Visible after filters',visible.length.toLocaleString());setMessage('',true)}
function field(parent,label,value,wide=false){const box=document.createElement('div');box.className='field'+(wide?' wide':'');const s=document.createElement('span'),b=document.createElement('b');s.textContent=label;b.textContent=show(value);box.append(s,b);parent.append(box)}
function disclosure(parent,label,value){const d=document.createElement('details'),s=document.createElement('summary'),pre=document.createElement('pre');s.textContent=label;pre.textContent=show(value);d.append(s,pre);parent.append(d)}
function drawDetail(r,pos,total){const root=$('#detail');root.replaceChildren();const top=document.createElement('div');top.className='detail-top';const title=document.createElement('b');title.textContent=`Card ${pos+1} of ${total} · PDF page ${show(r.pdf_page)} · grid ${show(r.grid_row)}, ${show(r.grid_col)}`;const nav=document.createElement('div');nav.className='nav';for(const [label,delta] of [['Previous',-1],['Next',1]]){const b=document.createElement('button');b.textContent=label;b.disabled=(pos+delta<0||pos+delta>=visible.length);b.onclick=()=>{selected=visible[pos+delta].i;render()};nav.append(b)}top.append(title,nav);root.append(top);
const content=document.createElement('div');content.className='content';const imcol=document.createElement('div'),box=document.createElement('div');box.className='image-box';if(r.image_url){const img=document.createElement('img');img.src=r.image_url;img.alt=`Voter card crop on page ${show(r.pdf_page)}`;img.onerror=()=>{box.replaceChildren();const e=document.createElement('div');e.className='empty-image';e.textContent='Card image not found. Keep the workbook’s _images folder beside the workbook.';box.append(e)};box.append(img)}else{const e=document.createElement('div');e.className='empty-image';e.textContent='No linked card image is available for this row. Re-run OCR with the updated code to save card crops.';box.append(e)}imcol.append(box);const cap=document.createElement('div');cap.className='caption';cap.textContent='Source card crop — compare the printed values with the extracted fields.';imcol.append(cap);
const data=document.createElement('div'),fields=document.createElement('div');fields.className='fields';for(const [label,key] of [['Serial number','serial_no'],['Serial status','serial_status'],['Voter ID','voter_id'],['Voter ID status','voter_id_status'],['Name','name'],['Relation','relation_type'],['Relative name','relation_name'],['House number','house_no'],['Age','age'],['Gender','gender'],['Serial candidates','serial_candidates'],['Voter ID candidates','voter_id_candidates'],['Gender candidates','gender_candidates']])field(fields,label,r[key]);data.append(fields);content.append(imcol,data);root.append(content);
const ev=document.createElement('div');ev.className='evidence';disclosure(ev,'Raw OCR text',r.raw_text);disclosure(ev,'Voter ID targeted OCR',r.voter_id_raw);disclosure(ev,'Gender targeted OCR',r.gender_raw);root.append(ev)}
$('#workbook').addEventListener('change',loadRecords);$('#search').addEventListener('input',render);$('#filter').addEventListener('change',render);$('#refresh').addEventListener('click',loadBooks);loadBooks();
</script></body></html>"""


class DashboardHandler(BaseHTTPRequestHandler):
    workbook_dir: Path

    def log_message(self, fmt, *args):
        print(f"[dashboard] {self.address_string()} {fmt % args}", flush=True)

    def _send(self, body: bytes, content_type: str, status: int = 200):
        self.send_response(status)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("X-Content-Type-Options", "nosniff")
        self.end_headers()
        self.wfile.write(body)

    def _json(self, data, status: int = 200):
        self._send(json.dumps(data, ensure_ascii=False).encode("utf-8"), "application/json; charset=utf-8", status)

    def _workbook(self, name: str) -> Path:
        if not name or Path(name).name != name or not name.lower().endswith("_ocr.xlsx"):
            raise ValueError("Choose a workbook ending in _ocr.xlsx")
        path = (self.workbook_dir / name).resolve()
        if path.parent != self.workbook_dir.resolve() or not path.is_file():
            raise FileNotFoundError("Workbook not found in the selected folder")
        return path

    def _rows(self, workbook: Path):
        wb = load_workbook(workbook, data_only=True, read_only=False)
        try:
            if "Voters" not in wb.sheetnames:
                raise ValueError("Workbook does not contain a Voters sheet")
            ws = wb["Voters"]
            headers = [cell.value for cell in ws[1]]
            result = []
            image_dir = (workbook.parent / f"{workbook.stem}_images").resolve()
            for row_index in range(2, ws.max_row + 1):
                values = {str(key): ws.cell(row_index, col).value for col, key in enumerate(headers, start=1) if key}
                if not any(value not in (None, "") for value in values.values()):
                    continue
                values["image_url"] = None
                if "source_image" in headers:
                    cell = ws.cell(row_index, headers.index("source_image") + 1)
                    link = cell.hyperlink.target if cell.hyperlink else None
                    if link:
                        image_path = Path(link)
                        if not image_path.is_absolute():
                            image_path = (workbook.parent / image_path).resolve()
                        else:
                            image_path = image_path.resolve()
                        try:
                            image_path.relative_to(image_dir)
                        except ValueError:
                            pass
                        else:
                            if image_path.is_file() and image_path.suffix.lower() == ".png":
                                values["image_url"] = f"/api/image?workbook={quote(workbook.name)}&row={row_index}"
                result.append(values)
            return result
        finally:
            wb.close()

    def do_GET(self):
        request = urlparse(self.path)
        params = parse_qs(request.query)
        try:
            if request.path == "/":
                return self._send(PAGE.encode("utf-8"), "text/html; charset=utf-8")
            if request.path == "/api/workbooks":
                items = []
                for path in sorted(self.workbook_dir.glob("*_ocr.xlsx")):
                    if path.name.startswith("~$"):
                        continue
                    rows = self._rows(path)
                    items.append({"name": path.name, "records": len(rows)})
                return self._json(items)
            if request.path == "/api/records":
                name = params.get("workbook", [""])[0]
                workbook = self._workbook(name)
                return self._json({"workbook": name, "records": self._rows(workbook)})
            if request.path == "/api/image":
                name = params.get("workbook", [""])[0]
                workbook = self._workbook(name)
                row = int(params.get("row", ["0"])[0])
                if row < 2:
                    raise ValueError("Invalid workbook row")
                wb = load_workbook(workbook, data_only=True, read_only=False)
                try:
                    ws = wb["Voters"]
                    headers = [cell.value for cell in ws[1]]
                    if "source_image" not in headers or row > ws.max_row:
                        raise FileNotFoundError("No image for this row")
                    link = ws.cell(row, headers.index("source_image") + 1).hyperlink
                    if not link or not link.target:
                        raise FileNotFoundError("No image link for this row")
                    image_path = Path(link.target)
                    if not image_path.is_absolute():
                        image_path = (workbook.parent / image_path).resolve()
                    else:
                        image_path = image_path.resolve()
                    image_dir = (workbook.parent / f"{workbook.stem}_images").resolve()
                    image_path.relative_to(image_dir)
                    if image_path.suffix.lower() != ".png" or not image_path.is_file():
                        raise FileNotFoundError("Card image file is missing")
                    return self._send(image_path.read_bytes(), "image/png")
                finally:
                    wb.close()
            return self._json({"error": "Not found"}, 404)
        except (ValueError, FileNotFoundError, KeyError) as exc:
            return self._json({"error": str(exc)}, 400)
        except Exception as exc:
            return self._json({"error": f"Could not read workbook: {exc}"}, 500)


def serve_dashboard(workbook_dir="data/output", host="127.0.0.1", port=8765):
    folder = Path(workbook_dir).resolve()
    if not folder.is_dir():
        raise FileNotFoundError(f"Workbook folder does not exist: {folder}")
    handler = type("ConfiguredDashboardHandler", (DashboardHandler,), {"workbook_dir": folder})
    server = ThreadingHTTPServer((host, port), handler)
    print(f"Dashboard: http://{host}:{port}  (workbooks: {folder})", flush=True)
    print("Press Ctrl+C to stop the local dashboard.", flush=True)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nDashboard stopped.", flush=True)
    finally:
        server.server_close()
